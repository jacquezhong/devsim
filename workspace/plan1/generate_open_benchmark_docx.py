#!/usr/bin/env python3
"""Generate a DOCX draft for the open DEVSIM reverse-recovery benchmark."""

from __future__ import annotations

import re
import sys
import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
PLAN_DIR = ROOT / "workspace" / "plan1"
TOOLS_DIR = ROOT / "tools"
sys.path.insert(0, str(TOOLS_DIR))

from create_equation import add_formula  # noqa: E402


DEFAULT_MD_PATH = PLAN_DIR / "paper_draft_open_benchmark.md"
DEFAULT_OUT_PATH = PLAN_DIR / "基于开源DEVSIM的二极管反向恢复可复现仿真基准.docx"
FIG_DIR = PLAN_DIR / "figures" / "benchmark"


FIGURES = {
    "fig1": (
        FIG_DIR / "fig_pub_dc_iv.png",
        "图1  基准 PIN 二极管 DC I-V 曲线与目标电流工作点",
    ),
    "fig2": (
        FIG_DIR / "fig_pub_lifetime_summary.png",
        "图2  载流子寿命对反向恢复电荷和存储移动电荷的影响",
    ),
    "fig3": (
        FIG_DIR / "fig_pub_recovery_waveforms.png",
        "图3  固定正向电压与固定目标电流协议下的反向恢复波形",
    ),
    "fig4": (
        FIG_DIR / "fig_pub_forward_current_sweep.png",
        "图4  目标正向电流密度对反向恢复面电荷和存储移动面电荷的影响",
    ),
    "fig5": (
        FIG_DIR / "fig_pub_time_step_convergence.png",
        "图5  固定目标电流协议下的时间步长收敛性",
    ),
    "fig6": (
        FIG_DIR / "fig_pub_mesh_convergence.png",
        "图6  固定正向电压协议下的网格密度收敛性",
    ),
}


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(size)
    run.font.bold = bold


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    style.font.size = Pt(10.5)
    style.paragraph_format.line_spacing = 1.5


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    return text


def add_text_with_inline_formula(paragraph, text: str, bold: bool = False) -> None:
    """Add text while converting simple $...$ inline formulas to OMML."""
    parts = re.split(r"(\$[^$]+\$)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("$") and part.endswith("$"):
            add_formula(paragraph, part[1:-1], display_mode=False)
            continue
        run = paragraph.add_run(clean_inline(part))
        set_run_font(run, bold=bold)


def add_paragraph(doc: Document, text: str, indent: bool = True) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.28)
    add_text_with_inline_formula(paragraph, text)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)


def add_figure(doc: Document, key: str) -> None:
    path, caption = FIGURES[key]
    if not path.exists():
        add_paragraph(doc, f"[缺少图片：{path}]", indent=False)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(5.8))
    add_caption(doc, caption)


def add_heading(doc: Document, line: str) -> None:
    level = len(line) - len(line.lstrip("#"))
    text = clean_inline(line[level:].strip())
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = 16
    elif level == 2:
        size = 14
    else:
        size = 12
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [clean_inline(part.strip()) for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
            rows.append(parts)
        index += 1
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text_with_inline_formula(paragraph, value, bold=(row_idx == 0))


def add_reference_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.first_line_indent = Inches(-0.2)
    paragraph.paragraph_format.line_spacing = 1.5
    add_text_with_inline_formula(paragraph, clean_inline(text), bold=False)


def should_skip(line: str) -> bool:
    return line.startswith(">") or line.startswith("```") or line.startswith("- `workspace/")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_MD_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT_PATH)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = args.input if args.input.is_absolute() else ROOT / args.input
    output_path = args.output if args.output.is_absolute() else ROOT / args.output

    lines = input_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_doc_defaults(doc)

    inserted: set[str] = set()
    in_code = False
    in_references = False
    index = 0

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code or not stripped or stripped.startswith(">"):
            index += 1
            continue

        if stripped.startswith("#"):
            add_heading(doc, stripped)
            in_references = "参考文献" in stripped
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            if rows and rows[0] and "寿命 tau" in rows[0][0] and "fig2" not in inserted:
                add_figure(doc, "fig2")
                inserted.add("fig2")
            if rows and rows[0] and "时间步长" in rows[0][0] and "fig5" not in inserted:
                add_figure(doc, "fig5")
                inserted.add("fig5")
            continue

        if stripped.startswith("$$"):
            formula_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("$$"):
                formula_lines.append(lines[index].strip())
                index += 1
            paragraph = doc.add_paragraph()
            add_formula(paragraph, " ".join(formula_lines), display_mode=True)
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style=None)
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.line_spacing = 1.5
            add_text_with_inline_formula(paragraph, stripped[2:])
            index += 1
            continue

        if stripped.startswith("[") and in_references:
            add_reference_paragraph(doc, stripped)
            index += 1
            continue

        if not should_skip(stripped):
            add_paragraph(doc, stripped)
            if re.search(r"图\s*1(给出|展示)", stripped) and "fig1" not in inserted:
                add_figure(doc, "fig1")
                inserted.add("fig1")
            if re.search(r"图\s*2", stripped) and "fig2" not in inserted:
                add_figure(doc, "fig2")
                inserted.add("fig2")
            if re.search(r"图\s*3(比较|展示)", stripped) and "fig3" not in inserted:
                add_figure(doc, "fig3")
                inserted.add("fig3")
            if re.search(r"图\s*4", stripped) and "fig4" not in inserted:
                add_figure(doc, "fig4")
                inserted.add("fig4")
            if re.search(r"图\s*5", stripped) and "fig5" not in inserted:
                add_figure(doc, "fig5")
                inserted.add("fig5")
            if (
                re.search(r"图\s*6", stripped) or "网格扫描结果显示" in stripped
            ) and "fig6" not in inserted:
                add_figure(doc, "fig6")
                inserted.add("fig6")

        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()
